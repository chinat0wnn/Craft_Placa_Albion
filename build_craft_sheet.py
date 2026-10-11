import json, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

OUT = "Planilha_Craft_Placas.xlsx"
CITIES = ["Thetford", "Martlock", "Bridgewatch", "Lymhurst", "Fort Sterling"]
PIECES = [("ARMOR", "Armadura", 16), ("HEAD", "Elmo", 8), ("SHOES", "Botas", 8)]
VARIANTS = [  # codigo, nome exibido, usa artefato
    ("SET1", "Soldado (Soldier)", False),
    ("SET2", "Cavaleiro (Knight)", False),
    ("SET3", "Guardião (Guardian)", False),
    ("UNDEAD", "Graveguard", True),
    ("HELL", "Demônio (Demon)", True),
    ("KEEPER", "Juiz (Judicator)", True),
    ("FEY", "Duskweaver", True),
    ("AVALON", "Valor", True),
]
ORE_QTY = {2: 1, 3: 2, 4: 2, 5: 3, 6: 4, 7: 5, 8: 5}
BAR_KEYS = ["T2.0", "T3.0"] + [f"T{t}.{e}" for t in range(4, 9) for e in range(5)]
ITEM_KEYS = [f"T{t}.{e}" for t in range(4, 9) for e in range(5)]
BAR_ROW = {k: 3 + i for i, k in enumerate(BAR_KEYS)}  # linha em Refino/Minerio/Barras

# ------------------------------------------------------------------
# Dados oficiais do jogo (ao-bin-dumps) - valida receitas e valores
# ------------------------------------------------------------------
game = json.load(open("items.json", encoding="utf-8"))["items"]
index = {}
def walk(o):
    if isinstance(o, dict):
        un = o.get("@uniquename")
        if un and ("craftingrequirements" in o or "@itemvalue" in o):
            index.setdefault(un, o)
        for v in o.values():
            walk(v)
    elif isinstance(o, list):
        for v in o:
            walk(v)
walk(game)

names = {}
for line in open("items.txt", encoding="utf-8"):
    m = re.match(r"\s*\d+:\s+(\S+)\s+:\s+(.*)$", line.rstrip())
    if m:
        names[m.group(1)] = m.group(2).strip()

def first_recipe(un):
    cr = index[un]["craftingrequirements"]
    cr = cr[0] if isinstance(cr, list) else cr
    res = cr["craftresource"]
    res = res if isinstance(res, list) else [res]
    return [(r["@uniquename"], int(r["@count"])) for r in res]

artifact = {}  # (codigo, slot, tier) -> (nome, IV)
for code, disp, uses_art in VARIANTS:
    for slot, pname, nbars in PIECES:
        for t in range(4, 9):
            un = f"T{t}_{slot}_PLATE_{code}"
            rec = first_recipe(un)
            assert rec[0] == (f"T{t}_METALBAR", nbars), (un, rec)
            if uses_art:
                art_id, cnt = rec[1]
                assert cnt == 1 and "ARTEFACT" in art_id, (un, rec)
                nm = re.sub(r"^\w+'s ", "", names.get(art_id, art_id))
                artifact[(code, slot, t)] = (nm, int(index[art_id]["@itemvalue"]))
            else:
                assert len(rec) == 1, (un, rec)
print("Receitas validadas contra os dados do jogo.")
for code, disp, _ in VARIANTS:
    print(" ", code, "->", names.get(f"T4_ARMOR_PLATE_{code}", names.get(f"T4_ARMOR_PLATE_{code}")))

# precos pre-existentes da planilha original (minerio e barras)
def load_old():
    ore, bars = {}, {}
    try:
        wb0 = openpyxl.load_workbook("planilha de refino de material.xlsx", data_only=True)
        ws0 = wb0["PesquisaRefino"]
        for r in range(3, 69):
            a = ws0.cell(r, 1).value
            if a is None or isinstance(a, str) and not re.match(r"^\d", a.strip()):
                continue
            try:
                key = f"T{float(a):.1f}"
            except ValueError:
                continue
            vals = [ws0.cell(r, c).value for c in range(2, 7)]
            vals = [v if isinstance(v, (int, float)) else None for v in vals]
            if not any(v is not None for v in vals):
                continue
            (ore if r <= 34 else bars)[key] = vals
    except Exception as e:
        print("Aviso: nao foi possivel ler precos antigos:", e)
    return ore, bars
OLD_ORE, OLD_BARS = load_old()

# ------------------------------------------------------------------
# Estilos
# ------------------------------------------------------------------
INPUT = PatternFill("solid", start_color="FFF2CC")
HEAD = PatternFill("solid", start_color="1F4E78")
HEAD2 = PatternFill("solid", start_color="7F6000")
CALC = PatternFill("solid", start_color="E2EFDA")
GREY = PatternFill("solid", start_color="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WHITE_B = Font(bold=True, color="FFFFFF")
MONEY = '#,##0;-#,##0;;@'      # esconde zeros
MONEY0 = '#,##0'
PCT = '0.0%'

def header(ws, row, labels, fill=HEAD, start_col=1, height=45):
    for i, lab in enumerate(labels):
        c = ws.cell(row, start_col + i, lab)
        c.font = WHITE_B
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")

wb = openpyxl.Workbook()

# ------------------------------------------------------------------
# CONFIG
# ------------------------------------------------------------------
cfg = wb.active
cfg.title = "Config"
title(cfg, "Configurações gerais", "Células amarelas = você preenche. Tudo na planilha obedece a esta tela.")
header(cfg, 3, ["Parâmetro", "Valor", "Como usar"], height=22)
cfg_rows = [
    ("Taxa de compra (ordem de compra)", 0.0, PCT, "TaxaCompra", "0% se comprar direto da oferta; 2,5% se criar ordem de compra."),
    ("Taxa total de venda", 0.065, PCT, "TaxaVenda", "6,5% com Premium (4% imposto + 2,5% ordem de venda); 10,5% sem Premium."),
    ("Retorno de recursos no REFINO", 0.367, PCT, "RetRefino", "Ex.: 36,7% (cidade com bônus) ou 53,9% com foco. Use o valor da sua situação."),
    ("Retorno de recursos no CRAFT", 0.248, PCT, "RetCraft", "Ex.: 24,8% (cidade com bônus de placas), 15,2% sem bônus, ~47,9% com foco."),
    ("Taxa da barraca - Refino (prata/100 nutrição)", 465, MONEY0, "FeeRefino", "Taxa mostrada na fundição. Usada em: Valor do item x 0,1125 x taxa / 100."),
    ("Taxa da barraca - Craft (prata/100 nutrição)", 465, MONEY0, "FeeCraft", "Taxa mostrada na forja de armaduras."),
]
for i, (lab, val, fmt, name, note) in enumerate(cfg_rows):
    r = 4 + i
    cfg.cell(r, 1, lab).border = BORDER
    c = cfg.cell(r, 2, val)
    c.fill = INPUT; c.number_format = fmt; c.border = BORDER
    cfg.cell(r, 3, note).border = BORDER
    wb.defined_names[name] = DefinedName(name, attr_text=f"Config!$B${r}")
cfg["A12"] = "Como usar a planilha"
cfg["A12"].font = Font(bold=True, size=12, color="1F4E78")
steps = [
    "1) Preencha os preços de COMPRA nas abas Minerio, Barras e Artefatos (5 cidades). Deixe em branco o que não encontrou.",
    "2) Aba Refino: a planilha compara refinar x comprar a barra pronta e escolhe a mais barata. Para forçar, mude a coluna 'Origem' (Automático / Comprar / Refinar).",
    "3) Aba Itens: preencha os preços de VENDA das peças (5 cidades + Caerleon + Mercado Negro). Ela calcula custo, lucro e margem de TODAS as combinações; use o filtro para ordenar.",
    "4) Aba Craft: escolha Variante, Peça, Tier (T4.0 ... T8.4) e a quantidade. Veja custo, melhor local de venda, lucro unitário e total.",
    "5) Aba Lista_Compras: mostra o que comprar (barras, minérios e artefatos), em qual cidade e quanto isso custa.",
    "Observações: artefatos NÃO retornam no craft (confirmado nos dados do jogo). A taxa da barraca não é devolvida pelo retorno de recursos.",
    "Quantidades da Lista_Compras são líquidas (já descontado o retorno de recursos esperado).",
]
for i, s in enumerate(steps):
    cfg.cell(13 + i, 1, s)
widths(cfg, {"A": 48, "B": 14, "C": 95})

# ------------------------------------------------------------------
# LISTAS
# ------------------------------------------------------------------
lst = wb.create_sheet("Listas")
header(lst, 1, ["Variante", "", "Peça", "", "Tier"], height=22)
for i, (_, disp, _) in enumerate(VARIANTS):
    lst.cell(2 + i, 1, disp)
for i, (_, p, _) in enumerate(PIECES):
    lst.cell(2 + i, 3, p)
for i, k in enumerate(ITEM_KEYS):
    lst.cell(2 + i, 5, k)
widths(lst, {"A": 28, "C": 14, "E": 10})

# ------------------------------------------------------------------
# COMPRA MINERIO / BARRAS  (5 cidades como na planilha original)
# ------------------------------------------------------------------
def price_sheet(name, ttl, sub, old):
    ws = wb.create_sheet(name)
    title(ws, ttl, sub)
    header(ws, 2, ["Tier"] + CITIES + ["Menor preço", "Cidade mais barata"], height=30)
    ws.row_dimensions[2].height = 30
    for k in BAR_KEYS:
        r = BAR_ROW[k]
        ws.cell(r, 1, k).font = Font(bold=True)
        vals = old.get(k, [None] * 5)
        for j in range(5):
            c = ws.cell(r, 2 + j, vals[j])
            c.fill = INPUT; c.number_format = MONEY0; c.border = BORDER
        g = ws.cell(r, 7, f'=IF(COUNT(B{r}:F{r})=0,"",MIN(B{r}:F{r}))')
        g.number_format = MONEY0; g.fill = CALC; g.border = BORDER; g.font = Font(bold=True)
        h = ws.cell(r, 8, f'=IF(G{r}="","",INDEX($B$2:$F$2,MATCH(G{r},B{r}:F{r},0)))')
        h.fill = CALC; h.border = BORDER
    widths(ws, {"A": 9, "B": 13, "C": 13, "D": 13, "E": 13, "F": 14, "G": 14, "H": 20})
    ws.freeze_panes = "B3"
    return ws

price_sheet("Minerio", "Preço de compra do minério bruto (Iron Ore etc.)",
            "Preencha o menor preço visto em cada cidade (valores iniciais vieram da sua planilha original).", OLD_ORE)
price_sheet("Barras", "Preço de compra das barras de metal prontas (Metal Bar)",
            "Usado para comparar 'comprar barra' x 'refinar você mesmo' na aba Refino.", OLD_BARS)

# ------------------------------------------------------------------
# COMPRA ARTEFATOS
# ------------------------------------------------------------------
art = wb.create_sheet("Artefatos")
title(art, "Preço de compra dos artefatos (1 por peça craftada - não retorna no craft)",
      "O artefato é o mesmo para qualquer encantamento (.0 a .4) do mesmo tier.")
header(art, 2, ["Chave", "Variante", "Peça", "Tier", "Artefato", "Valor do item"] + CITIES + ["Menor preço", "Cidade mais barata"], height=30)
ART_ROW = {}
r = 3
for code, disp, uses in VARIANTS:
    if not uses:
        continue
    for slot, pname, _ in PIECES:
        for t in range(4, 9):
            nm, iv = artifact[(code, slot, t)]
            ART_ROW[(disp, pname, t)] = r
            art.cell(r, 1, f"{disp}|{pname}|T{t}")
            art.cell(r, 2, disp); art.cell(r, 3, pname); art.cell(r, 4, f"T{t}")
            art.cell(r, 5, nm); art.cell(r, 6, iv).number_format = MONEY0
            for j in range(5):
                c = art.cell(r, 7 + j)
                c.fill = INPUT; c.number_format = MONEY0; c.border = BORDER
            g = art.cell(r, 12, f'=IF(COUNT(G{r}:K{r})=0,"",MIN(G{r}:K{r}))')
            g.number_format = MONEY0; g.fill = CALC; g.border = BORDER; g.font = Font(bold=True)
            h = art.cell(r, 13, f'=IF(L{r}="","",INDEX($G$2:$K$2,MATCH(L{r},G{r}:K{r},0)))')
            h.fill = CALC; h.border = BORDER
            r += 1
ART_LAST = r - 1
widths(art, {"A": 30, "B": 18, "C": 10, "D": 7, "E": 28, "F": 11, "G": 12, "H": 12, "I": 12, "J": 12, "K": 13, "L": 13, "M": 19})
art.freeze_panes = "G3"
art.auto_filter.ref = f"A2:M{ART_LAST}"

# ------------------------------------------------------------------
# REFINO
# ------------------------------------------------------------------
ref = wb.create_sheet("Refino")
title(ref, "Refino de barras: refinar x comprar (custo unitário real)",
      "Receita: minério (qtd por tier) + 1 barra do tier abaixo (T4.x usa a T3.0; T5+ usa o mesmo encantamento). Retorno de recursos só incide sobre materiais, não sobre a taxa.")
header(ref, 2, ["Barra", "Tier", "Encant.", "Valor do item", "Qtd minério", "Barra inferior",
                "Preço minério (c/ taxa)", "Custo ao REFINAR", "Preço barra pronta (c/ taxa)", "Cidade (barra pronta)",
                "Origem", "CUSTO EFETIVO da barra", "Decisão",
                "Demanda direta (Craft)", "Demanda por refino superior", "Demanda total",
                "Barras a COMPRAR", "Barras a REFINAR", "Minério necessário", "Custo compra barras", "Custo minério", "Cidade minério"], height=48)
consumers = {k: [] for k in BAR_KEYS}
lower = {}
for k in BAR_KEYS:
    t, e = int(k[1]), int(k[3])
    if t == 2:
        lw = None
    elif t == 3:
        lw = "T2.0"
    elif t == 4:
        lw = "T3.0"
    else:
        lw = f"T{t-1}.{e}"
    lower[k] = lw
    if lw:
        consumers[lw].append(k)
first, last = BAR_ROW[BAR_KEYS[0]], BAR_ROW[BAR_KEYS[-1]]
for k in BAR_KEYS:
    r = BAR_ROW[k]
    t, e = int(k[1]), int(k[3])
    lw = lower[k]
    ref.cell(r, 1, k).font = Font(bold=True)
    ref.cell(r, 2, t); ref.cell(r, 3, e)
    ref.cell(r, 4, f"=2^B{r}*2^C{r}").number_format = MONEY0
    ref.cell(r, 5, ORE_QTY[t])
    ref.cell(r, 6, lw if lw else "-")
    ref.cell(r, 7, f'=IF(Minerio!G{r}="","",Minerio!G{r}*(1+TaxaCompra))')
    fee = f"D{r}*0.1125*FeeRefino/100"
    if lw:
        lr = BAR_ROW[lw]
        ref.cell(r, 8, f'=IF(OR(G{r}="",L{lr}=""),"",(E{r}*G{r}+L{lr})*(1-RetRefino)+{fee})')
    else:
        ref.cell(r, 8, f'=IF(G{r}="","",E{r}*G{r}*(1-RetRefino)+{fee})')
    ref.cell(r, 9, f'=IF(Barras!G{r}="","",Barras!G{r}*(1+TaxaCompra))')
    ref.cell(r, 10, f'=IF(Barras!H{r}="","",Barras!H{r})')
    c = ref.cell(r, 11, "Automático"); c.fill = INPUT
    ref.cell(r, 12, (f'=IF(K{r}="Comprar",I{r},IF(K{r}="Refinar",H{r},IF(AND(H{r}="",I{r}=""),"",'
                     f'IF(H{r}="",I{r},IF(I{r}="",H{r},MIN(H{r},I{r}))))))'))
    ref.cell(r, 13, (f'=IF(L{r}="","",IF(K{r}="Comprar","Comprar",IF(K{r}="Refinar","Refinar",'
                     f'IF(H{r}="","Comprar",IF(I{r}="","Refinar",IF(H{r}<=I{r},"Refinar","Comprar"))))))'))
    ref.cell(r, 14, f"=SUMIF(Craft!$C$6:$C$35,$A{r},Craft!$S$6:$S$35)")
    if consumers[k]:
        ref.cell(r, 15, "=" + "+".join(f'IF(M{BAR_ROW[c_]}="Refinar",P{BAR_ROW[c_]}*(1-RetRefino),0)' for c_ in consumers[k]))
    else:
        ref.cell(r, 15, 0)
    ref.cell(r, 16, f"=N{r}+O{r}")
    ref.cell(r, 17, f'=IF(M{r}="Comprar",P{r},0)')
    ref.cell(r, 18, f'=IF(M{r}="Refinar",P{r},0)')
    ref.cell(r, 19, f"=R{r}*E{r}*(1-RetRefino)")
    ref.cell(r, 20, f'=IF(AND(Q{r}>0,ISNUMBER(I{r})),Q{r}*I{r},0)')
    ref.cell(r, 21, f'=IF(AND(S{r}>0,ISNUMBER(G{r})),S{r}*G{r},0)')
    ref.cell(r, 22, f'=IF(Minerio!H{r}="","",Minerio!H{r})')
    for col in range(1, 23):
        cell = ref.cell(r, col)
        cell.border = BORDER
        if col in (7, 8, 9, 12, 20, 21):
            cell.number_format = MONEY0
        if col in (14, 15, 16, 17, 18, 19):
            cell.number_format = '#,##0.0;-#,##0.0;;@'
        if col in (8, 12, 13):
            cell.fill = CALC
    ref.cell(r, 12).font = Font(bold=True, color="006100")
dv = DataValidation(type="list", formula1='"Automático,Comprar,Refinar"', allow_blank=False)
ref.add_data_validation(dv)
dv.add(f"K{first}:K{last}")
widths(ref, {"A": 8, "B": 6, "C": 8, "D": 10, "E": 9, "F": 10, "G": 13, "H": 13, "I": 14, "J": 14, "K": 12, "L": 15, "M": 11,
             "N": 12, "O": 14, "P": 11, "Q": 11, "R": 11, "S": 12, "T": 13, "U": 13, "V": 14})
ref.freeze_panes = "B3"

# ------------------------------------------------------------------
# ITENS (preços de venda + custo/lucro de todas as combinações)
# ------------------------------------------------------------------
itm = wb.create_sheet("Itens")
title(itm, "Preços de venda e lucro de todas as peças de placa",
      "Preencha o preço de venda por cidade, Caerleon e Mercado Negro. Use o filtro (setas do cabeçalho) para ordenar por lucro/margem.")
H = ["Chave", "Variante", "Peça", "Tier", "Tier nº", "Encant.", "Qtd barras"] + CITIES + ["Caerleon", "Mercado Negro"] + \
    ["Melhor preço de venda", "Melhor local", "Custo barras (líq. retorno)", "Artefato", "Cidade do artefato",
     "Valor do item (IV)", "Taxa do ferreiro", "CUSTO TOTAL", "Receita líquida (após taxa)", "LUCRO por item", "Margem"]
header(itm, 3, H, height=48)
ITEM_ROW = {}
r = 4
for code, disp, uses in VARIANTS:
    for slot, pname, nbars in PIECES:
        for k in ITEM_KEYS:
            t, e = int(k[1]), int(k[3])
            ITEM_ROW[(disp, pname, k)] = r
            br = BAR_ROW[k]
            itm.cell(r, 1, f"{disp}|{pname}|{k}")
            itm.cell(r, 2, disp); itm.cell(r, 3, pname); itm.cell(r, 4, k)
            itm.cell(r, 5, t); itm.cell(r, 6, e); itm.cell(r, 7, nbars)
            for col in range(8, 15):
                c = itm.cell(r, col)
                c.fill = INPUT; c.number_format = MONEY0
            itm.cell(r, 15, f'=IF(COUNT(H{r}:N{r})=0,"",MAX(H{r}:N{r}))').number_format = MONEY0
            itm.cell(r, 16, f'=IF(O{r}="","",INDEX($H$3:$N$3,MATCH(O{r},H{r}:N{r},0)))')
            itm.cell(r, 17, f'=IF(Refino!$L${br}="","",G{r}*Refino!$L${br}*(1-RetCraft))').number_format = MONEY0
            if uses:
                ar = ART_ROW[(disp, pname, t)]
                itm.cell(r, 18, f'=IF(Artefatos!$L${ar}="","",Artefatos!$L${ar}*(1+TaxaCompra))').number_format = MONEY0
                itm.cell(r, 19, f'=IF(Artefatos!$M${ar}="","",Artefatos!$M${ar})')
                itm.cell(r, 20, f"=G{r}*Refino!$D${br}+Artefatos!$F${ar}").number_format = MONEY0
            else:
                itm.cell(r, 18, 0)
                itm.cell(r, 19, "-")
                itm.cell(r, 20, f"=G{r}*Refino!$D${br}").number_format = MONEY0
            itm.cell(r, 21, f"=T{r}*0.1125*FeeCraft/100").number_format = MONEY0
            itm.cell(r, 22, f'=IF(OR(Q{r}="",R{r}=""),"",Q{r}+R{r}+U{r})').number_format = MONEY0
            itm.cell(r, 23, f'=IF(O{r}="","",O{r}*(1-TaxaVenda))').number_format = MONEY0
            itm.cell(r, 24, f'=IF(OR(V{r}="",W{r}=""),"",W{r}-V{r})').number_format = '#,##0;[Red]-#,##0'
            itm.cell(r, 25, f'=IF(X{r}="","",X{r}/V{r})').number_format = PCT
            for col in (15, 16, 17, 22, 23, 24, 25):
                itm.cell(r, col).fill = CALC
            r += 1
ITEM_LAST = r - 1
itm.auto_filter.ref = f"A3:Y{ITEM_LAST}"
itm.freeze_panes = "E4"
widths(itm, {"A": 30, "B": 18, "C": 9, "D": 7, "E": 6, "F": 7, "G": 8, "H": 11, "I": 11, "J": 12, "K": 11, "L": 12, "M": 11,
             "N": 13, "O": 13, "P": 14, "Q": 14, "R": 11, "S": 15, "T": 11, "U": 11, "V": 13, "W": 14, "X": 13, "Y": 9})
itm.conditional_formatting.add(f"X4:X{ITEM_LAST}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100", bold=True)))

# ------------------------------------------------------------------
# CRAFT (planejador)
# ------------------------------------------------------------------
cr = wb.create_sheet("Craft")
title(cr, "Planejador de craft - Placas (armadura, elmo, botas)",
      "Escolha nas células amarelas (listas suspensas). Cada linha = um tipo de peça que você quer craftar.")
cr["A3"] = "LUCRO TOTAL ESPERADO:"; cr["A3"].font = Font(bold=True)
cr["D3"] = "=SUM(Q6:Q35)"; cr["D3"].number_format = '#,##0;[Red]-#,##0'; cr["D3"].font = Font(bold=True, size=12)
cr["F3"] = "CUSTO TOTAL DO LOTE:"; cr["F3"].font = Font(bold=True)
cr["I3"] = "=SUM(R6:R35)"; cr["I3"].number_format = MONEY0; cr["I3"].font = Font(bold=True, size=12)
header(cr, 5, ["Variante", "Peça", "Tier", "Qtd a craftar", "Linha em Itens", "Barras por peça",
               "Custo barras (líq. retorno)", "Artefato (preço)", "Cidade do artefato", "Taxa do ferreiro",
               "CUSTO TOTAL / peça", "Melhor preço de venda", "Melhor local de venda", "Receita líquida / peça",
               "LUCRO / peça", "Margem", "LUCRO TOTAL", "Custo do lote", "Barras líquidas (após retorno)",
               "Barras brutas", "Artefatos a comprar"], height=48)
def idx(col, r):
    return f'=IF(ISNUMBER($E{r}),INDEX(Itens!${col}$4:${col}${ITEM_LAST},$E{r}),"")'
for r in range(6, 36):
    for col in range(1, 5):
        c = cr.cell(r, col); c.fill = INPUT; c.border = BORDER
    cr.cell(r, 4).number_format = MONEY0
    cr.cell(r, 5, f'=IF(OR(A{r}="",B{r}="",C{r}=""),"",IFERROR(MATCH(A{r}&"|"&B{r}&"|"&C{r},Itens!$A$4:$A${ITEM_LAST},0),"?"))')
    cr.cell(r, 6, idx("G", r))
    cr.cell(r, 7, idx("Q", r))
    cr.cell(r, 8, idx("R", r))
    cr.cell(r, 9, idx("S", r))
    cr.cell(r, 10, idx("U", r))
    cr.cell(r, 11, idx("V", r))
    cr.cell(r, 12, idx("O", r))
    cr.cell(r, 13, idx("P", r))
    cr.cell(r, 14, idx("W", r))
    cr.cell(r, 15, idx("X", r))
    cr.cell(r, 16, idx("Y", r))
    cr.cell(r, 17, f'=IF(AND(ISNUMBER(O{r}),ISNUMBER(D{r})),O{r}*D{r},0)')
    cr.cell(r, 18, f'=IF(AND(ISNUMBER(K{r}),ISNUMBER(D{r})),K{r}*D{r},0)')
    cr.cell(r, 19, f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(D{r})),D{r}*F{r}*(1-RetCraft),0)')
    cr.cell(r, 20, f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(D{r})),D{r}*F{r},0)')
    cr.cell(r, 21, f'=IF(AND(ISNUMBER(H{r}),ISNUMBER(D{r})),IF(H{r}>0,D{r},0),0)')
    for col in range(5, 22):
        c = cr.cell(r, col); c.border = BORDER
        if col in (7, 8, 10, 11, 12, 14, 18):
            c.number_format = MONEY0
        if col in (15, 17):
            c.number_format = '#,##0;[Red]-#,##0'
        if col == 16:
            c.number_format = PCT
        if col in (19,):
            c.number_format = '#,##0.0'
        if col in (11, 15, 17):
            c.fill = CALC
dvA = DataValidation(type="list", formula1="=Listas!$A$2:$A$9", allow_blank=True)
dvB = DataValidation(type="list", formula1="=Listas!$C$2:$C$4", allow_blank=True)
dvC = DataValidation(type="list", formula1="=Listas!$E$2:$E$26", allow_blank=True)
for dvx, col in ((dvA, "A"), (dvB, "B"), (dvC, "C")):
    cr.add_data_validation(dvx)
    dvx.add(f"{col}6:{col}35")
cr.conditional_formatting.add("O6:O35", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100", bold=True)))
cr.conditional_formatting.add("Q6:Q35", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100", bold=True)))
widths(cr, {"A": 19, "B": 10, "C": 7, "D": 10, "E": 9, "F": 9, "G": 13, "H": 11, "I": 14, "J": 11, "K": 13, "L": 13, "M": 14,
            "N": 13, "O": 12, "P": 9, "Q": 13, "R": 13, "S": 13, "T": 10, "U": 11})
cr.freeze_panes = "D6"

# ------------------------------------------------------------------
# LISTA DE COMPRAS
# ------------------------------------------------------------------
lc = wb.create_sheet("Lista_Compras")
title(lc, "Lista de compras para o que está planejado na aba Craft",
      "Quantidades líquidas (descontado o retorno de recursos esperado). Linhas com zero ficam vazias.")
lc["A3"] = "TOTAL A INVESTIR (barras + minério + artefatos):"; lc["A3"].font = Font(bold=True)
lc["F3"] = "=E37+K37+Q37"; lc["F3"].number_format = MONEY0; lc["F3"].font = Font(bold=True, size=12)
lc["H3"] = "Lucro esperado:"; lc["H3"].font = Font(bold=True)
lc["I3"] = "=Craft!D3"; lc["I3"].number_format = '#,##0;[Red]-#,##0'; lc["I3"].font = Font(bold=True, size=12)
header(lc, 5, ["BARRAS a comprar", "Qtd", "Cidade", "Preço unit.", "Total"], fill=HEAD, height=30)
header(lc, 5, ["MINÉRIO a comprar", "Qtd", "Cidade", "Preço unit.", "Total"], fill=HEAD2, start_col=7, height=30)
header(lc, 5, ["ARTEFATOS a comprar", "Qtd", "Cidade", "Preço unit.", "Total"], fill=PatternFill("solid", start_color="375623"), start_col=13, height=30)
for i, k in enumerate(BAR_KEYS):
    r = 6 + i
    rr = BAR_ROW[k]
    lc.cell(r, 1, k)
    lc.cell(r, 2, f"=Refino!Q{rr}").number_format = '#,##0.0;-#,##0.0;;@'
    lc.cell(r, 3, f'=IF(B{r}>0,Refino!J{rr},"")')
    lc.cell(r, 4, f'=IF(B{r}>0,Refino!I{rr},"")').number_format = MONEY
    lc.cell(r, 5, f"=Refino!T{rr}").number_format = MONEY
    lc.cell(r, 7, k)
    lc.cell(r, 8, f"=Refino!S{rr}").number_format = '#,##0.0;-#,##0.0;;@'
    lc.cell(r, 9, f'=IF(H{r}>0,Refino!V{rr},"")')
    lc.cell(r, 10, f'=IF(H{r}>0,Refino!G{rr},"")').number_format = MONEY
    lc.cell(r, 11, f"=Refino!U{rr}").number_format = MONEY
for i in range(30):
    r = 6 + i
    cr_r = 6 + i
    lc.cell(r, 13, f'=IF(Craft!U{cr_r}>0,Craft!A{cr_r}&" "&Craft!B{cr_r}&" "&Craft!C{cr_r},"")')
    lc.cell(r, 14, f"=Craft!U{cr_r}").number_format = MONEY
    lc.cell(r, 15, f'=IF(N{r}>0,Craft!I{cr_r},"")')
    lc.cell(r, 16, f'=IF(N{r}>0,Craft!H{cr_r},"")').number_format = MONEY
    lc.cell(r, 17, f'=IF(AND(N{r}>0,ISNUMBER(P{r})),N{r}*P{r},0)').number_format = MONEY
lc["A37"] = "Subtotal barras"; lc["E37"] = "=SUM(E6:E35)"
lc["G37"] = "Subtotal minério"; lc["K37"] = "=SUM(K6:K35)"
lc["M37"] = "Subtotal artefatos"; lc["Q37"] = "=SUM(Q6:Q35)"
for a in ("A37", "G37", "M37"):
    lc[a].font = Font(bold=True)
for a in ("E37", "K37", "Q37"):
    lc[a].font = Font(bold=True); lc[a].number_format = MONEY0
widths(lc, {"A": 20, "B": 9, "C": 14, "D": 11, "E": 13, "F": 3, "G": 20, "H": 9, "I": 14, "J": 11, "K": 13, "L": 3,
            "M": 40, "N": 9, "O": 14, "P": 11, "Q": 13})

# ordem das abas
order = ["Config", "Craft", "Lista_Compras", "Itens", "Refino", "Minerio", "Barras", "Artefatos", "Listas"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("Salvo:", OUT, "| itens:", ITEM_LAST - 3, "| artefatos:", ART_LAST - 2)
