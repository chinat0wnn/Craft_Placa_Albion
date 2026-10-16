"""Gera gamedata.json (receitas/artefatos oficiais) e um data.json inicial com precos da planilha antiga."""
import json, os, re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

PLATE_VARIANTS = [
    ("SET1", "Soldado (Soldier)", False, False),
    ("SET2", "Cavaleiro (Knight)", False, False),
    ("SET3", "Guardião (Guardian)", False, False),
    ("UNDEAD", "Graveguard", True, False),
    ("HELL", "Demônio (Demon)", True, False),
    ("KEEPER", "Juiz (Judicator)", True, False),
    ("FEY", "Duskweaver", True, False),
    ("AVALON", "Valor", True, False),
    ("ROYAL", "Real (Royal)", False, True),
]

BOW_VARIANTS = [
    ("2H_BOW", "Arco (Bow)", False, 0, 32),
    ("2H_WARBOW", "Arco de Guerra (Warbow)", False, 0, 32),
    ("2H_LONGBOW", "Arco Longo (Longbow)", False, 0, 32),
    ("2H_LONGBOW_UNDEAD", "Arco Sussurrante (Whispering)", True, 0, 32),
    ("2H_BOW_HELL", "Arco Lamenturiento (Wailing)", True, 0, 32),
    ("2H_BOW_KEEPER", "Arco de Badon (Badon)", True, 0, 32),
    ("2H_BOW_AVALON", "Perfurador de Névoa (Mistpiercer)", True, 0, 32),
]

CROSSBOW_VARIANTS = [
    ("2H_CROSSBOW", "Besta (Crossbow)", False, 12, 20),
    ("2H_CROSSBOWLARGE", "Besta Pesada (Heavy Crossbow)", False, 12, 20),
    ("MAIN_1HCROSSBOW", "Besta Leve (Light Crossbow 1M)", False, 8, 16),
    ("2H_REPEATINGCROSSBOW_UNDEAD", "Repetidora Lamenturienta (Weeping)", True, 12, 20),
    ("2H_DUALCROSSBOW_HELL", "Atiradores de Dardos (Boltcasters)", True, 12, 20),
    ("2H_CROSSBOWLARGE_MORGANA", "Besta de Cerco (Siegebow)", True, 12, 20),
    ("2H_CROSSBOW_CANNON_AVALON", "Moldador de Energia (Energy Shaper)", True, 12, 20),
]

PIECES = [
    ("ARMOR", "Armadura", 16, 0, "plate"),
    ("HEAD", "Elmo", 8, 0, "plate"),
    ("SHOES", "Botas", 8, 0, "plate"),
    ("BOW", "Arco", 0, 32, "bow"),
    ("CROSSBOW", "Besta", 12, 20, "crossbow"),
]
CITIES = ["Thetford", "Martlock", "Bridgewatch", "Lymhurst", "Fort Sterling"]

ROYAL_SIGILS = {
    4: {"id": "QUESTITEM_TOKEN_ROYAL_T4", "name": "Emblema Real do Adepto (T4)", "tier": 4},
    5: {"id": "QUESTITEM_TOKEN_ROYAL_T5", "name": "Emblema Real do Perito (T5)", "tier": 5},
    6: {"id": "QUESTITEM_TOKEN_ROYAL_T6", "name": "Emblema Real do Mestre (T6)", "tier": 6},
    7: {"id": "QUESTITEM_TOKEN_ROYAL_T7", "name": "Emblema Real do Grão-Mestre (T7)", "tier": 7},
    8: {"id": "QUESTITEM_TOKEN_ROYAL_T8", "name": "Emblema Real do Ancião (T8)", "tier": 8},
}

ROYAL_SIGIL_QTY = {
    "HEAD": {4: 2, 5: 4, 6: 8, 7: 8, 8: 8},
    "ARMOR": {4: 4, 5: 8, 6: 16, 7: 16, 8: 16},
    "SHOES": {4: 2, 5: 4, 6: 8, 7: 8, 8: 8},
}

game = json.load(open(os.path.join(ROOT, "items.json"), encoding="utf-8"))["items"]
index = {}
def walk(o):
    if isinstance(o, dict):
        un = o.get("@uniquename")
        if un and ("craftingrequirements" in o or "@itemvalue" in o):
            index.setdefault(un, o)
        for v in o.values():
            walk(v)
    elif isinstance(o, list):
        for v in o:
            walk(v)
walk(game)

names = {}
for line in open(os.path.join(ROOT, "items.txt"), encoding="utf-8"):
    m = re.match(r"\s*\d+:\s+(\S+)\s+:\s+(.*)$", line.rstrip())
    if m:
        names[m.group(1)] = m.group(2).strip()

def first_recipe(un):
    cr = index[un]["craftingrequirements"]
    cr = cr[0] if isinstance(cr, list) else cr
    res = cr["craftresource"]
    res = res if isinstance(res, list) else [res]
    return [(r["@uniquename"], int(r["@count"])) for r in res]

artifacts, item_names = {}, {}
all_variants = []

# 1. Placas
for code, disp, uses, is_royal in PLATE_VARIANTS:
    all_variants.append({
        "code": code,
        "name": disp,
        "artifact": uses,
        "royal": is_royal,
        "cat": "plate",
    })
    if is_royal:
        for slot, pname, _, _, pcat in PIECES:
            if pcat != "plate":
                continue
            for t in range(4, 9):
                for e in range(5):
                    suf = f"@{e}" if e else ""
                    un = f"T{t}_{slot}_PLATE_{code}{suf}"
                    item_names[un] = names.get(un, un)
        continue

    for slot, pname, nbars, _, pcat in PIECES:
        if pcat != "plate":
            continue
        for t in range(4, 9):
            un = f"T{t}_{slot}_PLATE_{code}"
            rec = first_recipe(un)
            assert rec[0] == (f"T{t}_METALBAR", nbars), (un, rec)
            for e in range(5):
                suf = f"@{e}" if e else ""
                item_names[f"{un}{suf}"] = names.get(f"{un}{suf}", names.get(un, un))
            if uses:
                art_id, cnt = rec[1]
                assert cnt == 1 and "ARTEFACT" in art_id
                artifacts[f"{code}_{slot}_{t}"] = {
                    "id": art_id,
                    "name": re.sub(r"^\w+'s ", "", names.get(art_id, art_id)),
                    "iv": int(index[art_id]["@itemvalue"]),
                }
            else:
                assert len(rec) == 1

# 2. Arcos
for code, disp, uses, nbars, nplanks in BOW_VARIANTS:
    all_variants.append({
        "code": code,
        "name": disp,
        "artifact": uses,
        "royal": False,
        "cat": "bow",
        "bars": nbars,
        "planks": nplanks,
    })
    slot = "BOW"
    for t in range(4, 9):
        un = f"T{t}_{code}"
        rec = first_recipe(un)
        assert rec[0] == (f"T{t}_PLANKS", nplanks), (un, rec)
        for e in range(5):
            suf = f"@{e}" if e else ""
            item_names[f"{un}{suf}"] = names.get(f"{un}{suf}", names.get(un, un))
        if uses:
            art_id, cnt = rec[1]
            assert cnt == 1 and "ARTEFACT" in art_id, (un, rec)
            artifacts[f"{code}_{slot}_{t}"] = {
                "id": art_id,
                "name": re.sub(r"^\w+'s ", "", names.get(art_id, art_id)),
                "iv": int(index[art_id]["@itemvalue"]),
            }
        else:
            assert len(rec) == 1, (un, rec)

# 3. Bestas
for code, disp, uses, nbars, nplanks in CROSSBOW_VARIANTS:
    all_variants.append({
        "code": code,
        "name": disp,
        "artifact": uses,
        "royal": False,
        "cat": "crossbow",
        "bars": nbars,
        "planks": nplanks,
    })
    slot = "CROSSBOW"
    for t in range(4, 9):
        un = f"T{t}_{code}"
        rec = first_recipe(un)
        assert rec[0] == (f"T{t}_PLANKS", nplanks), (un, rec)
        assert rec[1] == (f"T{t}_METALBAR", nbars), (un, rec)
        for e in range(5):
            suf = f"@{e}" if e else ""
            item_names[f"{un}{suf}"] = names.get(f"{un}{suf}", names.get(un, un))
        if uses:
            art_id, cnt = rec[2]
            assert cnt == 1 and "ARTEFACT" in art_id, (un, rec)
            artifacts[f"{code}_{slot}_{t}"] = {
                "id": art_id,
                "name": re.sub(r"^\w+'s ", "", names.get(art_id, art_id)),
                "iv": int(index[art_id]["@itemvalue"]),
            }
        else:
            assert len(rec) == 2, (un, rec)

for t, sig in ROYAL_SIGILS.items():
    item_names[sig["id"]] = names.get(sig["id"], sig["name"])

gd = {
    "cities": CITIES,
    "saleLocs": CITIES + ["Caerleon", "Black Market"],
    "categories": [
        {"id": "plate", "name": "🛡️ Placas"},
        {"id": "bow", "name": "🏹 Arcos"},
        {"id": "crossbow", "name": "🎯 Bestas"},
    ],
    "variants": all_variants,
    "pieces": [{"slot": s, "name": n, "bars": b, "planks": pl, "cat": c} for s, n, b, pl, c in PIECES],
    "artifacts": artifacts,
    "royal": {
        "sigils": {str(t): v for t, v in ROYAL_SIGILS.items()},
        "sigilQty": {s: {str(t): q for t, q in qt.items()} for s, qt in ROYAL_SIGIL_QTY.items()},
        "baseSets": ["SET1", "SET2", "SET3"],
    },
    "itemNames": item_names,
}
json.dump(gd, open(os.path.join(HERE, "gamedata.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("gamedata.json gerado:", len(artifacts), "artefatos,", len(ROYAL_SIGILS), "emblemas reais,", len(item_names), "nomes")

# ---- data.json inicial (nao sobrescreve se ja existir) ----
data_path = os.path.join(HERE, "data.json")
if not os.path.exists(data_path):
    prices = {}
    def bar_id(t, e): return f"T{t}_METALBAR" + (f"_LEVEL{e}@{e}" if e else "")
    def ore_id(t, e): return f"T{t}_ORE" + (f"_LEVEL{e}@{e}" if e else "")
    try:
        import openpyxl
        ws = openpyxl.load_workbook(os.path.join(ROOT, "planilha de refino de material.xlsx"), data_only=True)["PesquisaRefino"]
        for r in range(3, 69):
            a = ws.cell(r, 1).value
            try:
                t, e = (int(x) for x in f"{float(a):.1f}".split("."))
            except (TypeError, ValueError):
                continue
            vals = {CITIES[j]: ws.cell(r, 2 + j).value for j in range(5)}
            vals = {k: v for k, v in vals.items() if isinstance(v, (int, float)) and v > 0}
            if vals:
                prices[(ore_id if r <= 34 else bar_id)(t, e)] = vals
        print("Precos importados da planilha original:", list(prices))
    except Exception as ex:
        print("Sem precos iniciais:", ex)
    state = {
        "config": {"taxBuy": 0.0, "taxSell": 0.065, "retRefine": 0.367, "retCraft": 0.248,
                   "feeRefine": 465, "feeCraft": 465, "server": "west"},
        "prices": prices, "dates": {}, "origins": {}, "plan": [],
    }
    json.dump(state, open(data_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("data.json inicial criado")
